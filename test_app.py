import os
import re
import tempfile
import unittest
import io
from pathlib import Path

os.environ.setdefault('LOCAL_TEST_PASSWORD_HASH', 'test-hash')
os.environ.setdefault('APP_SECRET_KEY', 'unit-test-session-secret')

import app as application
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash


class LocalAppTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        application.DB = Path(self.temp_dir.name) / 'tests.sqlite3'
        application.app.config.update(TESTING=True, WTF_CSRF_ENABLED=True)
        os.environ['LOCAL_TEST_PASSWORD_HASH'] = generate_password_hash('local-test-password')
        application.init_db()
        self.client = application.app.test_client()

    def tearDown(self):
        self.temp_dir.cleanup()

    def csrf_token(self, response):
        match = re.search(rb'name="csrf_token" value="([^"]+)"', response.data)
        self.assertIsNotNone(match)
        return match.group(1).decode()

    def login(self):
        response = self.client.get('/login')
        token = self.csrf_token(response)
        return self.client.post(
            '/login',
            data={'csrf_token': token, 'password': 'local-test-password'},
            follow_redirects=True,
        )

    def test_authentication_and_csrf(self):
        self.assertEqual(self.client.get('/').status_code, 302)
        self.assertEqual(self.client.post('/logout').status_code, 400)
        response = self.login()
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Run local test batch', response.data)
        dashboard_token = self.csrf_token(response)
        logged_out = self.client.post('/logout', data={'csrf_token': dashboard_token})
        self.assertEqual(logged_out.status_code, 302)
        self.assertEqual(self.client.get('/').status_code, 302)

    def test_local_batch_persistence_and_excel_export(self):
        self.login()
        dashboard = self.client.get('/')
        token = self.csrf_token(dashboard)
        response = self.client.post(
            '/run-tests',
            data={'csrf_token': token, 'names': 'Tester One\nTester Two'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'SUBMITTED_LOCAL', response.data)
        with application.db() as con:
            rows = con.execute('SELECT * FROM submissions ORDER BY id').fetchall()
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['name'], 'Tester One')
        record_response = self.client.get(f"/record/{rows[0]['submission_key']}")
        self.assertEqual(record_response.status_code, 200)
        export = self.client.get('/export.xlsx')
        workbook = load_workbook(filename=io.BytesIO(export.data), data_only=True)
        self.assertEqual(workbook['Test Records'].max_row, 3)
        self.assertEqual(workbook['Test Records']['B2'].value, 'Tester One')

    def test_missing_names_and_missing_local_record(self):
        self.login()
        dashboard = self.client.get('/')
        token = self.csrf_token(dashboard)
        empty = self.client.post('/run-tests', data={'csrf_token': token, 'names': ' \n'})
        self.assertEqual(empty.status_code, 302)
        self.assertEqual(self.client.get('/record/does-not-exist').status_code, 404)

    def test_sample_form_rejects_missing_or_unlisted_options(self):
        self.login()
        dashboard = self.client.get('/')
        token = self.csrf_token(dashboard)
        form = {
            'csrf_token': token,
            'name': 'Local tester',
            'language': 'English',
            'initial_consent': 'yes',
            'final_consent': 'yes',
            'promoter': 'Promoter-1',
            'age_group': 'Age ≤ 50 years',
            'gender': 'Others',
            'swelling_both_legs': 'No',
            'worse_evening': 'No',
            'whole_body_swelling': 'No',
            'face_swelling_morning': 'No',
            'breathing_difficulty': 'No',
            'breathing_lying_down': 'No',
            'breathing_walking': 'No',
        }
        invalid = dict(form, promoter='Not an available option')
        self.assertEqual(self.client.post('/submit', data=invalid).status_code, 400)
        accepted = self.client.post('/submit', data=form)
        self.assertEqual(accepted.status_code, 200)
        self.assertIn(b'SUBMITTED_LOCAL', accepted.data)

    def test_excel_formula_values_are_escaped(self):
        self.assertEqual(application.excel_safe(' =1+1'), "' =1+1")


if __name__ == '__main__':
    unittest.main()
