import io
import tempfile
import unittest
from types import SimpleNamespace
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import threading
from urllib.parse import urlsplit
from pathlib import Path
from unittest.mock import patch
from copy import deepcopy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import app as application
import live_batches
from live_batches import FIELDS
from browser_agent import DryRunFailure, LocalBrowser, questionnaire_fingerprint
from consent import affirmative_consent
from playwright.sync_api import sync_playwright
from playwright.sync_api import Error as PlaywrightError, TimeoutError as PlaywrightTimeoutError


class LiveBatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.original_db = application.DB
        application.DB = Path(self.temp.name) / 'batches.sqlite3'
        self.original_csrf = application.app.config.get('WTF_CSRF_ENABLED')
        application.app.config['WTF_CSRF_ENABLED'] = False
        application.init_db()
        self.client = application.app.test_client()
        with self.client.session_transaction() as session:
            session['authenticated'] = True

    def tearDown(self):
        application.DB = self.original_db
        application.app.config['WTF_CSRF_ENABLED'] = self.original_csrf
        self.temp.cleanup()

    @staticmethod
    def workbook():
        book = Workbook()
        sheet = book.active
        sheet.title = 'Patients'
        headers = list(FIELDS)
        headers[0] = 'Person'
        sheet.append(headers)
        sheet.append(['Example', 'Promoter-24', 'Age< 50 years', 'Male', 'No', 'Yes', 'No', 'Yes', 'No', 'No', 'Yes', 'Anaemic', 'Liver disease', 'Yes'])
        sheet.append([])
        sheet.append(['Example 2', 'Promoter-24', 'Age< 50 years', 'Male', 'No', 'Yes', 'No', 'Yes', 'No', 'No', 'Yes', '', '', 'No'])
        book.create_sheet('Other')['A1'] = 'Preserve this sheet'
        out = io.BytesIO()
        book.save(out)
        out.seek(0)
        return out

    def upload(self):
        result = self.client.post('/live-entry/upload', data={'workbook': (self.workbook(), 'patients.xlsx')}, content_type='multipart/form-data')
        self.assertEqual(result.status_code, 302)
        return int(result.headers['Location'].rstrip('/').split('/')[-1])

    def test_upload_mapping_export_and_submission_gate(self):
        batch_id = self.upload()
        self.assertEqual(self.client.get(f'/live-entry/{batch_id}').status_code, 200)
        form = {field: field for field in FIELDS}
        form['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=form)
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,error FROM live_batch_rows ORDER BY row_number').fetchall()
        self.assertEqual([(r['row_number'], r['status']) for r in rows], [(2, 'PENDING'), (4, 'FAILED')])
        self.assertIn('Consent', rows[1]['error'])
        self.assertEqual(self.client.post(f'/live-entry/{batch_id}/submit').status_code, 403)
        response = self.client.get(f'/live-entry/{batch_id}/export.xlsx')
        book = load_workbook(io.BytesIO(response.data))
        sheet = book['Patients']
        self.assertEqual(sheet['A2'].value, 'Example')
        self.assertEqual(sheet['A4'].value, 'Example 2')
        self.assertEqual(sheet['O1'].value, 'Form ID')
        self.assertEqual(sheet['P1'].value, 'Status')
        self.assertEqual(sheet['P2'].value, 'PENDING')
        self.assertEqual(sheet['P4'].value, 'FAILED')
        self.assertIsNone(sheet['P3'].value)
        self.assertEqual(book['Other']['A1'].value, 'Preserve this sheet')

    def test_bad_headers_mapping_and_authentication(self):
        batch_id = self.upload()
        form = {field: 'Person' for field in FIELDS}
        self.client.post(f'/live-entry/{batch_id}/mapping', data=form)
        with application.db() as con:
            self.assertEqual(con.execute('SELECT header_mapping FROM live_batches WHERE id=?', (batch_id,)).fetchone()[0], '{}')
        anonymous = application.app.test_client()
        self.assertEqual(anonymous.get(f'/live-entry/{batch_id}/export.xlsx').status_code, 302)

    def test_consent_values_are_explicit(self):
        for value in ('Yes', '  yEs ', 'True', True, '1', 1, 'Tick', 'SELECTED'):
            with self.subTest(value=value):
                self.assertTrue(affirmative_consent(value))
        for value in (None, '', '   ', 'No', 'False', False, '0', 0, 'maybe', 'checked', 2):
            with self.subTest(value=value):
                self.assertFalse(affirmative_consent(value))

    def test_automatic_mapping_aliases_and_ambiguity(self):
        definition = live_batches.saved_mapping()
        headers = [field['label'] for field in definition['fields']]
        headers[0] = '  FULL   NAME  '
        headers[2] = 'Age'
        headers[3] = 'Gender'
        headers[8] = 'A. Do you have breathing difficulty.'
        headers[13] = 'I Consent'
        headers.extend(['Form ID', 'Status'])
        resolved, errors = live_batches.resolve_headers(headers, definition)
        self.assertEqual(errors, [])
        self.assertEqual(set(resolved), set(FIELDS))
        self.assertEqual(resolved['Breathing_Difficulty'], headers[8])
        self.assertEqual(definition['fields'][8]['label'], '6 A. Do you have breathing difficulty.')
        _, errors = live_batches.resolve_headers(headers + [' full name '], definition)
        self.assertTrue(any('Duplicate' in error or 'multiple' in error for error in errors))
        _, errors = live_batches.resolve_headers(headers[:2], definition)
        self.assertTrue(any('no matching header' in error for error in errors))

    def test_automatic_mapping_and_multiselect_values(self):
        book = load_workbook(self.workbook())
        sheet = book.active
        sheet['A1'] = 'Full Name'
        sheet['N2'] = 'Selected'
        sheet['L2'] = 'Anaemic\nDiabetes; High Blood pressure'
        sheet['M2'] = 'Liver disease, Kidney disease'
        out = io.BytesIO()
        book.save(out)
        out.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (out, 'auto.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        with application.db() as con:
            batch = con.execute('SELECT header_mapping FROM live_batches WHERE id=?', (batch_id,)).fetchone()
            row = con.execute('SELECT status,error FROM live_batch_rows WHERE batch_id=? AND row_number=2', (batch_id,)).fetchone()
        self.assertNotEqual(batch['header_mapping'], '{}')
        self.assertEqual(row['status'], 'PENDING', row['error'])

    def test_existing_report_columns_are_not_patient_answers_or_overwritten(self):
        book = load_workbook(self.workbook())
        sheet = book.active
        sheet['O1'] = 'Form ID'
        sheet['P1'] = 'Status'
        sheet['O2'] = 'Existing identifier'
        out = io.BytesIO()
        book.save(out)
        out.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (out, 'report.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        self.assertEqual(self.client.get(f'/live-entry/{batch_id}/export.xlsx').status_code, 400)

    def test_interrupted_submitting_becomes_unknown(self):
        batch_id = self.upload()
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='SUBMITTING' WHERE batch_id=? AND row_number=2", (batch_id,))
        application.init_db()
        with application.db() as con:
            status = con.execute('SELECT status FROM live_batch_rows WHERE batch_id=? AND row_number=2', (batch_id,)).fetchone()[0]
        self.assertEqual(status, 'UNKNOWN')
        self.assertEqual(self.client.post(f'/live-entry/{batch_id}/submit').status_code, 403)

    def test_saved_mapping_detects_structural_changes(self):
        mapping = live_batches.saved_mapping()
        fields = deepcopy(mapping['fields'])
        self.assertEqual(questionnaire_fingerprint(fields, mapping['fields']), mapping['fingerprint'])
        fields[2]['options'][0]['label'] = 'Changed'
        self.assertNotEqual(questionnaire_fingerprint(fields, mapping['fields']), mapping['fingerprint'])
        with self.assertRaisesRegex(ValueError, 'structure changed'):
            questionnaire_fingerprint(fields[:-1], mapping['fields'])

    def test_dry_run_processes_only_valid_row_with_mocked_browser(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        batch = live_batches.get_batch(batch_id)
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        with patch.object(live_batches.browser_worker, 'call', return_value={'status': 'VALIDATED'}) as worker:
            live_batches._run_dry_batch(batch_id, batch['workbook'], columns, live_batches.saved_mapping())
        self.assertEqual(worker.call_count, 1)
        self.assertEqual(worker.call_args.args[2]['Full Name'], 'Example')
        with application.db() as con:
            rows = con.execute('SELECT row_number,status FROM live_batch_rows ORDER BY row_number').fetchall()
        self.assertEqual([(r['row_number'], r['status']) for r in rows], [(2, 'VALIDATED'), (4, 'FAILED')])

    def test_bad_options_and_multiple_choice(self):
        original = self.workbook()
        book = load_workbook(original)
        sheet = book.active
        sheet['B2'] = 'Invalid promoter'
        sheet['L2'] = 'Anaemic; Diabetes'
        sheet['L4'] = 'None'
        sheet['M4'] = 'None'
        sheet['N4'] = 'Yes'
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (output, 'test.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,error FROM live_batch_rows ORDER BY row_number').fetchall()
        self.assertEqual(rows[0]['status'], 'FAILED')
        self.assertIn('Promoter', rows[0]['error'])
        self.assertNotIn('Clinical_Conditions', rows[0]['error'])
        self.assertEqual(rows[1]['status'], 'PENDING')

    def test_export_only_places_verified_ids_on_submitted_rows(self):
        batch_id = self.upload()
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='SUBMITTED',form_id='website-123' WHERE batch_id=? AND row_number=2", (batch_id,))
            con.execute("UPDATE live_batch_rows SET status='UNKNOWN',form_id='unverified' WHERE batch_id=? AND row_number=4", (batch_id,))
        response = self.client.get(f'/live-entry/{batch_id}/export.xlsx')
        sheet = load_workbook(io.BytesIO(response.data)).active
        self.assertEqual(sheet['O2'].value, 'website-123')
        self.assertIsNone(sheet['O4'].value)
        self.assertEqual(sheet['P4'].value, 'UNKNOWN')

    def test_browser_failure_does_not_submit_or_leak_answers(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        batch = live_batches.get_batch(batch_id)
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        with patch.object(live_batches.browser_worker, 'call', side_effect=RuntimeError('sensitive detail')):
            live_batches._run_dry_batch(batch_id, batch['workbook'], columns, live_batches.saved_mapping())
        with application.db() as con:
            row = con.execute('SELECT status,error FROM live_batch_rows WHERE row_number=2').fetchone()
        self.assertEqual(row['status'], 'FAILED')
        self.assertNotIn('sensitive detail', row['error'])

    def test_browser_failure_does_not_stop_other_pending_rows(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='PENDING' WHERE batch_id=? AND row_number=4", (batch_id,))
        batch = live_batches.get_batch(batch_id)
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        failure = DryRunFailure('Browser launch: Chromium executable is missing.', systemic=True)
        with patch.object(live_batches.browser_worker, 'call', side_effect=failure) as worker:
            live_batches._run_dry_batch(batch_id, batch['workbook'], columns, live_batches.saved_mapping())
        self.assertEqual(worker.call_count, 2)
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,error FROM live_batch_rows ORDER BY row_number').fetchall()
        self.assertEqual([(r['row_number'], r['status']) for r in rows], [(2, 'FAILED'), (4, 'FAILED')])
        self.assertIn('Chromium executable is missing', rows[0]['error'])

    def test_row_specific_error_is_recorded_and_next_row_continues(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='PENDING' WHERE batch_id=? AND row_number=4", (batch_id,))
        batch = live_batches.get_batch(batch_id)
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        failure = DryRunFailure('Questionnaire answer mapping: Age: option is missing, invalid, or ambiguous.')
        with patch.object(live_batches.browser_worker, 'call', side_effect=[failure, {'status': 'VALIDATED'}]):
            live_batches._run_dry_batch(batch_id, batch['workbook'], columns, live_batches.saved_mapping())
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,error FROM live_batch_rows ORDER BY row_number').fetchall()
        self.assertEqual([(r['row_number'], r['status']) for r in rows], [(2, 'FAILED'), (4, 'VALIDATED')])
        self.assertIn('Age:', rows[0]['error'])

    def test_three_rows_continue_after_middle_error_and_persist_immediately(self):
        book = load_workbook(self.workbook())
        sheet = book.active
        sheet['L4'], sheet['M4'], sheet['N4'] = 'None', 'None', 'Yes'
        for col in range(1, 15):
            sheet.cell(5, col, sheet.cell(2, col).value)
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (output, 'fixture.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        batch = live_batches.get_batch(batch_id)
        calls = []
        def browser_call(*args):
            calls.append(1)
            if len(calls) == 2:
                with application.db() as con:
                    state = con.execute('SELECT row_number,status FROM live_batch_rows ORDER BY row_number').fetchall()
                self.assertEqual([(r['row_number'], r['status']) for r in state], [(2, 'VALIDATED'), (4, 'VALIDATING'), (5, 'PENDING')])
                raise DryRunFailure('Age: invalid answer.')
            return {'status': 'VALIDATED'}
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        with patch.object(live_batches.browser_worker, 'call', side_effect=browser_call):
            live_batches._run_dry_batch(batch_id, batch['workbook'], columns, live_batches.saved_mapping())
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,error FROM live_batch_rows ORDER BY row_number').fetchall()
        self.assertEqual([(r['row_number'], r['status']) for r in rows], [(2, 'VALIDATED'), (4, 'FAILED'), (5, 'VALIDATED')])
        self.assertEqual(rows[1]['error'], 'Age: invalid answer.')
        progress = self.client.get(f'/live-entry/{batch_id}/progress').get_json()
        self.assertEqual((progress['total'], progress['processed'], progress['counts']['VALIDATED'], progress['counts']['FAILED']), (3, 3, 2, 1))
        report = load_workbook(io.BytesIO(self.client.get(f'/live-entry/{batch_id}/export.xlsx').data)).active
        self.assertEqual(report['Q1'].value, 'Failure reason')
        self.assertEqual(report['P4'].value, 'FAILED')
        self.assertEqual(report['Q4'].value, 'Age: invalid answer.')

    def test_invalid_workbook_row_between_valid_rows_does_not_block_validation(self):
        book = load_workbook(self.workbook())
        sheet = book.active
        for col in range(1, 15):
            sheet.cell(5, col, sheet.cell(2, col).value)
        sheet['D4'], sheet['L4'], sheet['M4'], sheet['N4'] = 'Invalid gender', 'None', 'None', 'Yes'
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (output, 'invalid-middle.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        batch = live_batches.get_batch(batch_id)
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        with patch.object(live_batches.browser_worker, 'call', return_value={'status':'VALIDATED'}) as worker:
            live_batches._run_dry_batch(batch_id, batch['workbook'], columns, live_batches.saved_mapping())
        self.assertEqual(worker.call_count, 2)
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,error FROM live_batch_rows ORDER BY row_number').fetchall()
        self.assertEqual([(r['row_number'],r['status']) for r in rows], [(2,'VALIDATED'),(4,'FAILED'),(5,'VALIDATED')])
        self.assertIn('Gender:', rows[1]['error'])

    def test_retry_excludes_attempted_and_uncertain_rows(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='FAILED',attempted_submit=1 WHERE batch_id=? AND row_number=2", (batch_id,))
            con.execute("UPDATE live_batch_rows SET status='UNKNOWN' WHERE batch_id=? AND row_number=4", (batch_id,))
        with patch.object(live_batches, 'start_dry_run') as start:
            self.client.post(f'/live-entry/{batch_id}/retry-failed')
        start.assert_not_called()
        with application.db() as con:
            self.assertEqual([(r['row_number'],r['status']) for r in con.execute('SELECT row_number,status FROM live_batch_rows ORDER BY row_number')], [(2,'FAILED'), (4,'UNKNOWN')])

    def test_retry_selects_only_eligible_failed_rows(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='FAILED',error='Browser error' WHERE batch_id=? AND row_number=2", (batch_id,))
            con.execute("UPDATE live_batch_rows SET status='UNKNOWN' WHERE batch_id=? AND row_number=4", (batch_id,))
        with patch.object(live_batches, 'start_dry_run') as start:
            self.client.post(f'/live-entry/{batch_id}/retry-failed')
        self.assertEqual(start.call_args.args[2], {2})
        with application.db() as con:
            self.assertEqual([(r['row_number'],r['status']) for r in con.execute('SELECT row_number,status FROM live_batch_rows ORDER BY row_number')], [(2,'PENDING'), (4,'UNKNOWN')])

    def test_browser_launch_error_diagnostic_does_not_expose_answers(self):
        agent = LocalBrowser()
        with patch.object(agent, 'open', side_effect=PlaywrightError("BrowserType.launch: Executable doesn't exist")), patch.object(agent, 'stop'):
            with self.assertRaises(DryRunFailure) as captured:
                agent.dry_run_questionnaire(live_batches.saved_mapping(), {'Consent': 'Selected', 'Full Name': 'private'})
        self.assertTrue(captured.exception.systemic)
        self.assertIn('Chromium executable is missing', str(captured.exception))
        self.assertNotIn('private', str(captured.exception))

    def test_timeout_diagnostic_has_stage_and_no_answers(self):
        agent = LocalBrowser()
        with patch.object(agent, 'open', side_effect=PlaywrightTimeoutError('private patient information')), patch.object(agent, 'stop'):
            with self.assertRaises(DryRunFailure) as captured:
                agent.dry_run_questionnaire(live_batches.saved_mapping(), {'Consent': 'Selected'})
        self.assertIn('Browser launch or page load: TimeoutError:', str(captured.exception))
        self.assertNotIn('private patient information', str(captured.exception))

    def test_rendered_manual_mapping_options_are_separate(self):
        batch_id = self.upload()
        html = self.client.get(f'/live-entry/{batch_id}').get_data(as_text=True)
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            page.set_content(html)
            if not page.locator('details').evaluate('(element) => element.open'):
                page.locator('details summary').click()
            control = page.locator('select[name="Full Name"]')
            self.assertTrue(control.is_visible())
            labels = control.locator('option').all_text_contents()
            self.assertEqual(labels, ['Select column', 'Person', *FIELDS[1:]])
            self.assertEqual(control.locator('option').count(), 15)
            self.assertEqual(control.locator('option[value="Person"]').count(), 1)
            browser.close()

    def test_submission_readiness_tracks_validated_and_failed_rows(self):
        batch_id = self.upload()
        path = f'/live-entry/{batch_id}'
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            def serve(route):
                target = urlsplit(route.request.url).path
                response = self.client.get(target)
                route.fulfill(status=response.status_code, body=response.get_data(),
                              content_type=response.content_type or 'text/html')
            page.route('**/*', serve)
            page.goto('http://dashboard.test' + path)
            button = page.locator('#submit-entries')
            page.wait_for_function("document.querySelector('#submission-reason').textContent.includes('0 validated candidates')")
            self.assertTrue(button.is_disabled())
            self.assertIn('No validated rows', page.locator('#submission-reason').inner_text())
            with application.db() as con:
                con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=?", (batch_id,))
            page.wait_for_function("document.querySelector('#submission-reason').textContent.includes('2 validated candidates')", timeout=8000)
            self.assertTrue(button.is_enabled())
            self.assertIn('Review the exact answers', page.locator('#submission-reason').inner_text())
            with application.db() as con:
                con.execute("UPDATE live_batch_rows SET status='FAILED' WHERE batch_id=? AND row_number=4", (batch_id,))
            page.wait_for_function("document.querySelector('#submission-reason').textContent.includes('1 validated candidates')", timeout=8000)
            self.assertTrue(button.is_enabled())
            self.assertIn('Failed: 1', page.locator('#progress-summary').inner_text())
            with application.db() as con:
                con.execute("UPDATE live_batch_rows SET attempted_submit=1 WHERE batch_id=? AND row_number=2", (batch_id,))
            page.wait_for_function("document.querySelector('#submission-reason').textContent.includes('0 validated candidates')", timeout=8000)
            response = self.client.post(path + '/submit')
            self.assertEqual(response.status_code, 403)
            self.assertIn(b'Review every answer', response.data)
            browser.close()

    def test_submission_attempts_are_durable_and_uncertain_rows_not_retried(self):
        book = load_workbook(self.workbook())
        sheet = book.active
        sheet['L4'], sheet['M4'], sheet['N4'] = 'None', 'None', 'Yes'
        for col in range(1, 15):
            sheet.cell(5, col, sheet.cell(2, col).value)
        data = io.BytesIO()
        book.save(data)
        data.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (data, 'fixture.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        headers = {field: field for field in FIELDS}
        headers['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=headers)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=?", (batch_id,))
        batch = live_batches.get_batch(batch_id)
        calls = []
        def submit_mock(*args, **kwargs):
            calls.append(1)
            with application.db() as con:
                attempt = con.execute("SELECT state FROM live_submission_attempts WHERE batch_id=? AND state='SUBMITTING'", (batch_id,)).fetchall()
            self.assertTrue(attempt, 'Attempt must be committed before the browser is invoked')
            if len(calls) == 1:
                return {'confirmation_verified': True, 'evidence_source': 'fixture', 'form_id': 'FORM-123'}
            raise TimeoutError('simulated browser crash')
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        with patch.object(live_batches, 'VERIFIED_CONFIRMATION_SOURCES', frozenset({'fixture'})), patch.object(live_batches.browser_worker, 'call', side_effect=submit_mock):
            live_batches._run_submission_batch(batch_id, batch['workbook'], headers, live_batches.saved_mapping(), {2, 4, 5}, operator_confirmed=True)
        self.assertEqual(len(calls), 3)
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,form_id,attempted_submit FROM live_batch_rows WHERE batch_id=? ORDER BY row_number', (batch_id,)).fetchall()
            attempts = con.execute('SELECT row_number,state FROM live_submission_attempts WHERE batch_id=? ORDER BY row_number', (batch_id,)).fetchall()
        self.assertEqual([(r['row_number'], r['status']) for r in rows], [(2, 'SUBMITTED'), (4, 'UNKNOWN'), (5, 'UNKNOWN')])
        self.assertEqual(rows[0]['form_id'], 'FORM-123')
        self.assertIsNone(rows[1]['form_id'])
        self.assertEqual([(r['row_number'], r['state']) for r in attempts], [(2, 'SUBMITTED'), (4, 'UNKNOWN'), (5, 'UNKNOWN')])
        with patch.object(live_batches, 'start_dry_run') as retry:
            self.client.post(f'/live-entry/{batch_id}/retry-failed')
        retry.assert_not_called()
        self.assertEqual(self.client.post(f'/live-entry/{batch_id}/submit').status_code, 403)

    def test_unverified_response_cannot_confirm_or_supply_a_form_id(self):
        batch_id = self.upload()
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=? AND row_number=2", (batch_id,))
        answers = {'Consent': 'Yes', 'Full Name': 'Fixture'}
        self.assertTrue(live_batches._claim_submission(batch_id, 2, answers))
        result = {'confirmation_verified': True, 'evidence_source': 'unverified', 'form_id': 'made-up'}
        self.assertEqual(live_batches._record_submission_result(batch_id, 2, result), 'UNKNOWN')
        with application.db() as con:
            row = con.execute('SELECT status,form_id FROM live_batch_rows WHERE batch_id=? AND row_number=2', (batch_id,)).fetchone()
        self.assertEqual(row['status'], 'UNKNOWN')
        self.assertIsNone(row['form_id'])

    def test_pending_attempt_becomes_unknown_on_restart(self):
        batch_id = self.upload()
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=? AND row_number=2", (batch_id,))
        self.assertTrue(live_batches._claim_submission(batch_id, 2, {'Consent':'Yes','Full Name':'Fixture'}))
        application.init_db()
        with application.db() as con:
            row = con.execute('SELECT status,attempted_submit FROM live_batch_rows WHERE batch_id=? AND row_number=2', (batch_id,)).fetchone()
            attempt = con.execute('SELECT state FROM live_submission_attempts WHERE batch_id=? AND row_number=2', (batch_id,)).fetchone()
        self.assertEqual((row['status'], row['attempted_submit'], attempt['state']), ('UNKNOWN', 1, 'UNKNOWN'))

    def test_identical_rows_across_batches_are_not_blocked(self):
        first, second = self.upload(), self.upload()
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE row_number=2")
        answers = {'Consent':'Selected', 'Full Name':'Fixture', 'Gender':'Male'}
        self.assertTrue(live_batches._claim_submission(first, 2, answers))
        live_batches._record_submission_result(first, 2, {})
        self.assertTrue(live_batches._claim_submission(second, 2, answers))
        with application.db() as con:
            row = con.execute('SELECT status,error FROM live_batch_rows WHERE batch_id=? AND row_number=2', (second,)).fetchone()
            attempts = con.execute('SELECT COUNT(*) FROM live_submission_attempts').fetchone()[0]
        self.assertEqual(row['status'], 'SUBMITTING')
        self.assertEqual(attempts, 2)

    def test_review_and_confirm_submits_only_validated_rows_once(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=? AND row_number=2", (batch_id,))
            con.execute("UPDATE live_batch_rows SET status='FAILED' WHERE batch_id=? AND row_number=4", (batch_id,))
        self.assertEqual(self.client.post(f'/live-entry/{batch_id}/submit').status_code, 403)
        review = self.client.get(f'/live-entry/{batch_id}/review')
        self.assertEqual(review.status_code, 200)
        self.assertEqual(review.headers['Cache-Control'], 'no-store')
        self.assertIn(b'Example', review.data)
        self.assertIn(b'Promoter-24', review.data)
        with self.client.session_transaction() as session:
            token = session['live_submission_review']['token']
        def thread_factory(target, args, kwargs, daemon):
            return SimpleNamespace(start=lambda: target(*args, **kwargs))
        result = {'confirmation_verified': True, 'evidence_source': 'score-page'}
        with patch.object(live_batches.threading, 'Thread', side_effect=thread_factory), patch.object(live_batches.browser_worker, 'call', return_value=result) as worker:
            response = self.client.post(f'/live-entry/{batch_id}/submit', data={'review_token':token, 'authorized':'yes'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(worker.call_count, 1)
        with application.db() as con:
            rows = con.execute('SELECT row_number,status,form_id FROM live_batch_rows WHERE batch_id=? ORDER BY row_number', (batch_id,)).fetchall()
        self.assertEqual([(r['row_number'],r['status']) for r in rows], [(2,'SUBMITTED'),(4,'FAILED')])
        self.assertIsNone(rows[0]['form_id'])
        report = load_workbook(io.BytesIO(self.client.get(f'/live-entry/{batch_id}/export.xlsx').data)).active
        self.assertIsNone(report['O2'].value)
        self.assertEqual(report['P2'].value, 'SUBMITTED')
        self.assertEqual(report['P4'].value, 'FAILED')
        self.assertEqual(self.client.post(f'/live-entry/{batch_id}/submit', data={'review_token':token, 'authorized':'yes'}).status_code, 403)

    def test_review_rejects_unconfirmed_or_changed_batch(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=? AND row_number=2", (batch_id,))
            con.execute("UPDATE live_batch_rows SET status='FAILED' WHERE batch_id=? AND row_number=4", (batch_id,))
        self.client.get(f'/live-entry/{batch_id}/review')
        with self.client.session_transaction() as session:
            token = session['live_submission_review']['token']
        self.assertEqual(self.client.post(f'/live-entry/{batch_id}/submit', data={'review_token':token}).status_code, 403)
        self.client.get(f'/live-entry/{batch_id}/review')
        with self.client.session_transaction() as session:
            token = session['live_submission_review']['token']
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=? AND row_number=4", (batch_id,))
        self.assertEqual(self.client.post(f'/live-entry/{batch_id}/submit', data={'review_token':token,'authorized':'yes'}).status_code, 409)

    def test_submission_rejects_concurrent_browser_job(self):
        batch_id = self.upload()
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=? AND row_number=2", (batch_id,))
            con.execute("UPDATE live_batch_rows SET status='FAILED' WHERE batch_id=? AND row_number=4", (batch_id,))
        self.client.get(f'/live-entry/{batch_id}/review')
        with self.client.session_transaction() as session:
            token = session['live_submission_review']['token']
        self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
        try:
            with patch.object(live_batches.browser_worker, 'call') as worker:
                response = self.client.post(f'/live-entry/{batch_id}/submit', data={'review_token':token,'authorized':'yes'})
            self.assertEqual(response.status_code, 409)
            worker.assert_not_called()
        finally:
            live_batches.batch_lock.release()

    def test_score_page_requires_post_and_visible_score(self):
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            url = 'http://score.test/form?lan_id=3'
            def serve(route):
                if route.request.method == 'POST':
                    route.fulfill(status=200, content_type='text/html', body='<h1>Total Score -6</h1>')
                else:
                    route.fulfill(status=200, content_type='text/html', body='<form method="post" action="#"><button type="submit">Submit</button></form>')
            page.route('**/*', serve)
            page.goto(url)
            agent = LocalBrowser()
            agent.page = page
            self.assertEqual(agent.submit_to_score_page(url, timeout_ms=3000), {'confirmation_verified': True, 'evidence_source': 'score-page'})
            browser.close()

    def test_score_page_does_not_accept_http_failure_or_unrelated_page(self):
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            url = 'http://score.test/form?lan_id=3'
            def serve(route):
                if route.request.method == 'POST':
                    route.fulfill(status=500, content_type='text/html', body='<h1>Total Score -6</h1>')
                else:
                    route.fulfill(status=200, content_type='text/html', body='<form method="post"><button type="submit">Submit</button></form>')
            page.route('**/*', serve)
            page.goto(url)
            agent = LocalBrowser()
            agent.page = page
            self.assertEqual(agent.submit_to_score_page(url, timeout_ms=3000), {'confirmation_verified': False, 'uncertainty_reason': 'Website returned HTTP 500; outcome unverified.'})
            browser.close()

    def test_success_http_without_score_is_unverified(self):
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            url = 'http://score.test/form?lan_id=3'
            requests = []
            def serve(route):
                if route.request.method == 'POST':
                    requests.append('POST')
                    route.fulfill(status=200, content_type='text/html', body='<h1>Unrelated page</h1>')
                else:
                    route.fulfill(status=200, content_type='text/html', body='<form method="post"><button type="submit">Submit</button></form>')
            page.route('**/*', serve)
            page.goto(url)
            agent = LocalBrowser()
            agent.page = page
            with self.assertRaises(PlaywrightTimeoutError):
                agent.submit_to_score_page(url, timeout_ms=1000)
            self.assertEqual(requests, ['POST'])
            browser.close()

    def test_post_redirect_to_score_page_is_recognized(self):
        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass
            def do_GET(self):
                body = b'<h1>Total Score -6</h1>' if self.path == '/survey/score' else b'<form method="post"><button type="submit">Submit</button></form>'
                self.send_response(200)
                self.send_header('Content-Type', 'text/html')
                self.send_header('Content-Length', str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            def do_POST(self):
                self.send_response(303)
                self.send_header('Location', '/survey/score')
                self.end_headers()
        server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                page = browser.new_page()
                url = f'http://127.0.0.1:{server.server_port}/survey/form?lan_id=3'
                page.goto(url)
                agent = LocalBrowser()
                agent.page = page
                self.assertEqual(agent.submit_to_score_page(url, timeout_ms=5000), {'confirmation_verified': True, 'evidence_source': 'score-page'})
                browser.close()
        finally:
            server.shutdown()
            thread.join(timeout=5)
            server.server_close()


    def test_submission_readiness_waits_for_pending_or_active_validation(self):
        rows = [dict(status='VALIDATED',attempted_submit=0,form_id=None), dict(status='FAILED',attempted_submit=0,form_id=None)]
        self.assertTrue(live_batches.submission_readiness(rows)['enabled'])
        self.assertFalse(live_batches.submission_readiness(rows, running=True)['enabled'])
        rows[1]['status'] = 'PENDING'
        self.assertFalse(live_batches.submission_readiness(rows)['enabled'])

    def test_submission_worker_requires_confirmation_and_verified_source(self):
        batch_id = self.upload()
        with application.db() as con:
            con.execute("UPDATE live_batch_rows SET status='VALIDATED' WHERE batch_id=? AND row_number=2", (batch_id,))
        batch = live_batches.get_batch(batch_id)
        mapping = {field: field for field in FIELDS}
        mapping['Full Name'] = 'Person'
        for confirmed in (False, True):
            self.assertTrue(live_batches.batch_lock.acquire(blocking=False))
            with patch.object(live_batches, 'VERIFIED_CONFIRMATION_SOURCES', frozenset()), patch.object(live_batches.browser_worker, 'call') as worker:
                live_batches._run_submission_batch(batch_id, batch['workbook'], mapping, live_batches.saved_mapping(), {2}, operator_confirmed=confirmed)
            worker.assert_not_called()
        with application.db() as con:
            row = con.execute('SELECT status,attempted_submit FROM live_batch_rows WHERE batch_id=? AND row_number=2', (batch_id,)).fetchone()
        self.assertEqual((row['status'],row['attempted_submit']), ('VALIDATED',0))

    def test_formatting_gaps_header_detection_and_export_alignment(self):
        book = load_workbook(self.workbook())
        sheet = book.active
        sheet.insert_rows(1, 2)
        sheet['A1'] = 'Patient workbook'
        sheet['XFD1048576'].fill = PatternFill(fill_type='solid', fgColor='FF0000')
        out = io.BytesIO()
        book.save(out)
        out.seek(0)
        result = self.client.post('/live-entry/upload', data={'workbook': (out, 'styled.xlsx')}, content_type='multipart/form-data')
        self.assertEqual(result.status_code, 302)
        batch_id = int(result.headers['Location'].rstrip('/').split('/')[-1])
        response = self.client.get(f'/live-entry/{batch_id}')
        self.assertIn(b'Detected header row: 3', response.data)
        self.assertIn(b'Actual non-empty data rows: 2', response.data)
        self.assertIn(b'Populated columns: 14', response.data)
        with application.db() as con:
            numbers = [row[0] for row in con.execute('SELECT row_number FROM live_batch_rows ORDER BY row_number')]
        self.assertEqual(numbers, [4, 6])
        exported = load_workbook(io.BytesIO(self.client.get(f'/live-entry/{batch_id}/export.xlsx').data))
        self.assertEqual(exported.active['A1'].value, 'Patient workbook')
        self.assertEqual(exported.active['O3'].value, 'Form ID')
        self.assertEqual(exported.active['P4'].value, 'PENDING')
        self.assertIsNone(exported.active['P5'].value)
        self.assertEqual(exported.active['P6'].value, 'PENDING')
        self.assertEqual(exported.active['XFD1048576'].fill.fgColor.rgb, '00FF0000')

    def test_formula_cell_counts_as_data_and_is_not_silently_accepted(self):
        book = load_workbook(self.workbook())
        book.active['A10'] = '=IF(1=1,"Example","")'
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (output, 'formula.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        columns = {field: field for field in FIELDS}
        columns['Full Name'] = 'Person'
        self.client.post(f'/live-entry/{batch_id}/mapping', data=columns)
        with application.db() as con:
            row = con.execute('SELECT status,error FROM live_batch_rows WHERE row_number=10').fetchone()
        self.assertEqual(row['status'], 'FAILED')
        self.assertIn('formulas are not accepted', row['error'])
        exported = load_workbook(io.BytesIO(self.client.get(f'/live-entry/{batch_id}/export.xlsx').data))
        self.assertEqual(exported.active['A10'].value, '=IF(1=1,"Example","")')

    def test_sparse_far_data_is_retained_and_counted(self):
        book = load_workbook(self.workbook())
        book.active['Z1048500'] = 'Actual data'
        output = io.BytesIO()
        book.save(output)
        output.seek(0)
        response = self.client.post('/live-entry/upload', data={'workbook': (output, 'sparse.xlsx')}, content_type='multipart/form-data')
        batch_id = int(response.headers['Location'].rstrip('/').split('/')[-1])
        detail = self.client.get(f'/live-entry/{batch_id}')
        self.assertIn(b'Actual non-empty data rows: 3', detail.data)
        self.assertIn(b'Populated columns: 15', detail.data)
        exported = load_workbook(io.BytesIO(self.client.get(f'/live-entry/{batch_id}/export.xlsx').data))
        self.assertEqual(exported.active['Z1048500'].value, 'Actual data')
        self.assertEqual(exported.active['AA1'].value, 'Form ID')
        self.assertEqual(exported.active['AB1048500'].value, 'PENDING')

    def test_rejects_genuinely_oversized_populated_data(self):
        book = Workbook()
        sheet = book.active
        sheet['A1'] = 'Header'
        for number in range(2, 20003):
            sheet.cell(number, 1, 'record')
        out = io.BytesIO()
        book.save(out)
        with self.assertRaisesRegex(ValueError, '20,000 populated data rows'):
            live_batches.open_workbook(out.getvalue())
        book = Workbook()
        for column in range(1, 258):
            book.active.cell(1, column, f'Header {column}')
        out = io.BytesIO()
        book.save(out)
        with self.assertRaisesRegex(ValueError, '256 populated columns'):
            live_batches.open_workbook(out.getvalue())


if __name__ == '__main__':
    unittest.main()
