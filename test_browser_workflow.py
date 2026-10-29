import os
import re
import socket
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from werkzeug.security import generate_password_hash
from werkzeug.serving import make_server
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
import io
import json

os.environ.setdefault('APP_SECRET_KEY', 'browser-test-session-secret')
os.environ.setdefault('LOCAL_TEST_PASSWORD_HASH', 'unused-test-hash')
os.environ['LOCAL_TEST_HEADLESS'] = '1'

import app as application
from browser_agent import LocalBrowser, validate_local_url, validate_target_url


class BrowserWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        application.DB = Path(cls.temp_dir.name) / 'browser-tests.sqlite3'
        application.app.config.update(TESTING=True, WTF_CSRF_ENABLED=True)
        os.environ['LOCAL_TEST_PASSWORD_HASH'] = generate_password_hash('browser-local-password')
        application.init_db()
        with socket.socket() as sock:
            sock.bind(('127.0.0.1', 0))
            cls.port = sock.getsockname()[1]
        cls.server = make_server('127.0.0.1', cls.port, application.app, threaded=True)
        cls.server_thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.server_thread.start()

    @classmethod
    def tearDownClass(cls):
        application.browser_worker.call('stop')
        cls.server.shutdown()
        cls.server_thread.join(timeout=5)
        cls.server.server_close()
        cls.temp_dir.cleanup()
        os.environ.pop('LOCAL_TEST_HEADLESS', None)

    def setUp(self):
        application.browser_worker.call('stop')
        application.active_browser_run_id = None
        with application.db() as con:
            con.execute('DELETE FROM browser_runs')
            con.execute('DELETE FROM submissions')
            con.execute("UPDATE app_settings SET setting_value='6' WHERE setting_key='daily_limit'")
        self.client = application.app.test_client()
        login_page = self.client.get('/login')
        token = self._csrf(login_page.data)
        self.client.post(
            '/login',
            data={'csrf_token': token, 'password': 'browser-local-password'},
        )
        self.token = self._csrf(self.client.get('/automation').data)
        self.target_url = f'http://127.0.0.1:{self.port}/survey'

    @staticmethod
    def _csrf(content):
        match = re.search(rb'name="csrf_token" value="([^"]+)"', content)
        if not match:
            raise AssertionError('The page did not contain a CSRF token.')
        return match.group(1).decode()

    def _post(self, path, data):
        return self.client.post(
            path,
            json=data,
            headers={'X-CSRFToken': self.token},
        )

    def _open_and_detect(self):
        forbidden = self._post('/browser/open', {'url': 'https://example.com/form'})
        self.assertEqual(forbidden.status_code, 400)
        opened = self._post('/browser/open', {'url': self.target_url})
        self.assertEqual(opened.status_code, 200, opened.get_data(as_text=True))
        inspected = self._post('/browser/detect', {})
        self.assertEqual(inspected.status_code, 200, inspected.get_data(as_text=True))
        return inspected.get_json()['fields']

    def _rules(self, fields, name_action='value'):
        rules = {}
        for field in fields:
            if field['name'] in {'initial_consent', 'final_consent'}:
                action, value = 'manual', None
            elif field['name'] == 'name':
                action, value = name_action, ('Local Browser Tester' if name_action == 'value' else None)
            elif field['options']:
                action, value = 'random', None
            else:
                action, value = 'skip', None
            rules[field['key']] = {'action': action, 'value': value}
        return rules

    def test_loopback_url_validation(self):
        self.assertEqual(validate_local_url('http://127.0.0.1:8765/survey'), 'http://127.0.0.1:8765/survey')
        self.assertTrue(validate_local_url('http://localhost:8765/survey').startswith('http://localhost:8765/survey'))
        for url in (
            'https://example.com/form',
            'file:///C:/secret',
            'http://127.0.0.1.evil.test/',
            'http://127.0.0.1:8765/survey?access_token=secret',
        ):
            with self.assertRaises(ValueError):
                validate_local_url(url)

    def test_live_url_scope(self):
        good = 'https://seeedemaseekhelp.com/Ma_Vaishno_Devi/form.php?lan_id=3'
        self.assertEqual(validate_target_url(good, allow_live=True), good)
        for url in ('https://example.com/', 'http://seeedemaseekhelp.com/Ma_Vaishno_Devi/',
                    'https://seeedemaseekhelp.com/Ma_Vaishno_Devi_evil/',
                    'https://seeedemaseekhelp.com:8443/Ma_Vaishno_Devi/'):
            with self.assertRaises(ValueError):
                validate_target_url(url, allow_live=True)

    def test_live_browser_blocks_private_network_resources(self):
        class Request:
            url = 'http://192.168.1.2/resource'
            def is_navigation_request(self): return False
        class Route:
            request = Request()
            aborted = False
            def abort(self): self.aborted = True
            def continue_(self): raise AssertionError('Private resource was allowed')
        route = Route()
        LocalBrowser._restrict_live_request(route)
        self.assertTrue(route.aborted)

    def test_nearby_question_labels_and_option_readback(self):
        self._open_and_detect()
        application.browser_worker.call('set_page_content', '''<form>
          <div><label>Full Name</label><input name="person" required></div>
          <label>Age</label><br><input type="radio" name="age" value="a" required><label>Under 50</label>
          <input type="radio" name="age" value="b" required><label>Over 50</label><br>
          <label>Conditions</label><br><ul><input type="checkbox" name="conditions" value="c"><label>Anaemic</label></ul>
          <ul><input type="checkbox" name="conditions" value="d"><label>Other</label></ul>
          <label for="p">Promoter</label><select id="p" name="promoter" required><option value="">Choose</option><option value="v24">Promoter-24</option></select>
          <label><input type="checkbox" name="consent" required> I Consent</label></form>''')
        fields = self._post('/browser/detect', {}).get_json()['fields']
        by_name = {f['name']: f for f in fields}
        self.assertEqual(by_name['person']['label'], 'Full Name')
        self.assertEqual(by_name['age']['label'], 'Age')
        self.assertEqual([o['label'] for o in by_name['age']['options']], ['Under 50', 'Over 50'])
        self.assertEqual(by_name['conditions']['label'], 'Conditions')
        self.assertEqual([o['label'] for o in by_name['conditions']['options']], ['Anaemic', 'Other'])
        application.browser_worker.call('check_selector', '[name="consent"]')
        answers = {'person':'Ayush', 'age':'Under 50', 'conditions':['Anaemic'], 'promoter':'Promoter-24'}
        rules = {f['key']: ({'action':'manual'} if f['name']=='consent' else
                 {'action':'value', 'value':answers[f['name']], **({'match':'label'} if f['options'] else {})}) for f in fields}
        plan = application.browser_worker.call('build_plan', rules, '')
        self.assertEqual(next(p for p in plan if p['label']=='Promoter')['expected'], 'v24')
        results = application.browser_worker.call('apply_and_validate', plan)
        self.assertTrue(all(r['result']=='VALIDATED' for r in results), results)

    def test_ambiguous_option_and_missing_required_are_rejected(self):
        self._open_and_detect()
        application.browser_worker.call('set_page_content', '<form><label for="x">Choice</label><select id="x" required><option value="a">Same</option><option value="b">Same</option></select><input name="new" required></form>')
        fields = self._post('/browser/detect', {}).get_json()['fields']
        rules = {fields[0]['key']:{'action':'value','match':'label','value':'Same'}, fields[1]['key']:{'action':'manual'}}
        with self.assertRaisesRegex(ValueError, 'ambiguous'):
            application.browser_worker.call('build_plan', rules, '')
        rules[fields[0]['key']] = {'action':'manual'}
        rules[fields[1]['key']] = {'action':'skip'}
        with self.assertRaisesRegex(ValueError, 'required field cannot be skipped'):
            application.browser_worker.call('build_plan', rules, '')

    def test_live_dashboard_preview_and_submission_block(self):
        url = 'https://seeedemaseekhelp.com/Ma_Vaishno_Devi/form.php?lan_id=3'
        with self.client.session_transaction() as session:
            session['browser_target_url'] = 'https://seeedemaseekhelp.com/Ma_Vaishno_Devi/'
        def worker(method, *args, **kwargs):
            if method == 'status':
                return {'open': True, 'fields_stale': False, 'url': url}
            if method == 'build_plan':
                return [{'action': 'manual', 'label': 'Question', 'expected': None}]
            raise AssertionError(method)
        with patch.object(application.browser_worker, 'call', side_effect=worker):
            preview = self._post('/browser/preview', {'name': 'Fixture', 'rules': {}})
            self.assertEqual(preview.status_code, 200, preview.get_data(as_text=True))
            with application.db() as con:
                con.execute("UPDATE browser_runs SET status='VALIDATED' WHERE id=?", (preview.get_json()['run_id'],))
            submitted = self._post('/browser/submit', {})
            self.assertEqual(submitted.status_code, 403)

    def test_multistep_fixture_dataset_without_submission(self):
        self._open_and_detect()
        questions = [('Age','Age< 50 years'), ('Gender','Male'), ('Both legs','No'),
                     ('Worse evening','Yes'), ('All over','No'), ('Face morning','Yes'),
                     ('Breathing','No'), ('Lying down','No'), ('Walking','Yes')]
        html = '<form><label>Full Name<input name="name" required></label><label>Promoter<select name="promoter" required><option value="">Choose</option><option value="opaque-promoter">Promoter-24</option></select></label>'
        for index, (question, answer) in enumerate(questions):
            html += f'<fieldset><legend>{question}</legend><label><input type="radio" name="q{index}" value="opaque-{index}" required>{answer}</label></fieldset>'
        html += '<fieldset><legend>Conditions</legend><label><input type="checkbox" name="conditions" value="opaque-condition">Anaemic</label></fieldset><fieldset><legend>Disorders</legend><label><input type="checkbox" name="disorders" value="opaque-disorder">Liver disease</label></fieldset><label><input type="checkbox" name="consent" required>I Consent</label><button type="submit">Submit</button></form>'
        application.browser_worker.call('set_page_content', '<select name="language"><option>English</option></select><label><input type="checkbox" name="initial_consent" required>I Consent</label><button id="next">Continue</button>')
        self._post('/browser/detect', {})
        application.browser_worker.call('check_selector', '[name="initial_consent"]')
        application.browser_worker.call('evaluate_page', '() => { document.querySelector("#next").onclick = () => { document.body.innerHTML = ' + json.dumps(html) + '; }; document.querySelector("#next").click(); }')
        self.assertTrue(self.client.get('/browser/status').get_json()['fields_stale'])
        fields = self._post('/browser/detect', {}).get_json()['fields']
        self.assertEqual(len(fields), 14)
        application.browser_worker.call('check_selector', '[name="consent"]')
        wanted = {'name':'Ayush', 'promoter':'Promoter-24', 'conditions':['Anaemic'], 'disorders':['Liver disease']}
        wanted.update({f'q{i}':answer for i, (_, answer) in enumerate(questions)})
        rules = {f['key']: ({'action':'manual'} if f['name']=='consent' else
                 {'action':'value','value':wanted[f['name']], **({'match':'label'} if f['options'] else {})}) for f in fields}
        results = application.browser_worker.call('apply_and_validate', application.browser_worker.call('build_plan', rules, ''))
        self.assertTrue(all(r['result']=='VALIDATED' for r in results), results)

    def test_dashboard_cookie_is_not_forwarded_to_live_host(self):
        with patch.object(application.browser_worker, 'call', return_value={'open': True}) as call:
            response = self._post('/browser/open', {'url': 'https://seeedemaseekhelp.com/Ma_Vaishno_Devi/'})
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(call.call_args.kwargs['session_cookie'])
        self.assertTrue(call.call_args.kwargs['allow_live'])

    def test_local_open_keeps_network_restriction(self):
        with patch.object(application.browser_worker, 'call', return_value={'open': True}) as call:
            response = self._post('/browser/open', {'url': self.target_url})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(call.call_args.kwargs['allow_live'])

    def test_browser_context_blocks_non_loopback_navigation(self):
        self._open_and_detect()
        class FakeRequest:
            url = 'https://example.com/'

        class FakeRoute:
            request = FakeRequest()
            aborted = False

            def abort(self):
                self.aborted = True

            def continue_(self):
                raise AssertionError('The external request was unexpectedly allowed.')

        route = FakeRoute()
        LocalBrowser._restrict_request(route)
        self.assertTrue(route.aborted)

    def test_live_field_detection_required_validation_submission_and_history(self):
        fields = self._open_and_detect()
        by_name = {field['name']: field for field in fields}
        self.assertIn('Full Name', by_name['name']['label'])
        self.assertEqual(by_name['promoter']['type'], 'select')
        self.assertTrue(any(option['value'] == 'Promoter-1' for option in by_name['promoter']['options']))
        self.assertEqual(by_name['gender']['type'], 'radio')
        self.assertEqual(by_name['conditions']['type'], 'checkbox')
        self.assertTrue(any(option['value'] == 'Hindi' for option in by_name['language']['options']))

        application.browser_worker.call('check_selector', '[name="initial_consent"]')
        application.browser_worker.call('check_selector', '[name="final_consent"]')

        consent_rules = self._rules(fields)
        consent_field = next(field for field in fields if field['name'] == 'initial_consent')
        consent_rules[consent_field['key']] = {'action': 'value', 'value': ['yes']}
        consent_preview = self._post('/browser/preview', {
            'name': 'Consent Must Be Manual',
            'rules': consent_rules,
        })
        self.assertEqual(consent_preview.status_code, 400)

        invalid_preview = self._post('/browser/preview', {
            'name': 'Blank Required Name',
            'seed': 'reproducible',
            'rules': self._rules(fields, name_action='manual'),
        })
        self.assertEqual(invalid_preview.status_code, 200, invalid_preview.get_data(as_text=True))
        invalid_fill = self._post('/browser/fill', {})
        self.assertEqual(invalid_fill.status_code, 200)
        self.assertEqual(invalid_fill.get_json()['status'], 'VALIDATION_FAILED')
        self.assertTrue(any(item['label'] == 'Full Name' and item['result'] == 'VALIDATION_FAILED'
                            for item in invalid_fill.get_json()['results']))

        valid_preview = self._post('/browser/preview', {
            'name': 'Valid Browser Tester',
            'seed': 'reproducible',
            'rules': self._rules(fields),
        })
        self.assertEqual(valid_preview.status_code, 200, valid_preview.get_data(as_text=True))
        preview_plan = valid_preview.get_json()['plan']
        name_plan = next(item for item in preview_plan if item['label'] == 'Full Name')
        self.assertEqual(name_plan['expected'], 'Local Browser Tester')
        self.assertTrue(any(item['action'] == 'random' for item in preview_plan))
        reproducible_plan = application.browser_worker.call(
            'build_plan', self._rules(fields), 'reproducible',
        )
        self.assertEqual(
            [item['expected'] for item in preview_plan],
            [item['expected'] for item in reproducible_plan],
        )
        valid_fill = self._post('/browser/fill', {})
        self.assertEqual(valid_fill.status_code, 200)
        self.assertEqual(valid_fill.get_json()['status'], 'VALIDATED')
        self.assertTrue(all(item['result'] in {'VALIDATED', 'SKIPPED'}
                            for item in valid_fill.get_json()['results']))

        duplicate = self._post('/browser/preview', {
            'name': 'valid browser tester',
            'seed': 'reproducible',
            'rules': self._rules(fields),
        })
        self.assertEqual(duplicate.status_code, 409)

        submitted = self._post('/browser/submit', {})
        self.assertEqual(submitted.status_code, 200, submitted.get_data(as_text=True))
        result = submitted.get_json()
        self.assertEqual(result['status'], 'SUBMITTED_CONFIRMED')
        self.assertTrue(result['record_id'])
        self.assertTrue(result['record_link'].startswith('http://127.0.0.1:'))
        with application.db() as con:
            run = con.execute(
                'SELECT status, website_record_id FROM browser_runs WHERE id=?',
                (result['run_id'],),
            ).fetchone()
            local_record = con.execute(
                'SELECT status FROM submissions WHERE id=?',
                (int(result['record_id']),),
            ).fetchone()
        self.assertEqual(run['status'], 'SUBMITTED_CONFIRMED')
        self.assertEqual(local_record['status'], 'SUBMITTED_LOCAL')
        self.assertEqual(self.client.get(
            '/history?test_name=Valid+Browser+Tester&status=SUBMITTED_CONFIRMED&sort=status'
        ).status_code, 200)
        self.assertEqual(self.client.get(f"/browser-record/{result['run_id']}").status_code, 200)
        workbook = load_workbook(io.BytesIO(self.client.get('/export.xlsx').data), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ['Test Records', 'Field Results', 'Detected Fields', 'Errors'],
        )
        self.assertEqual(workbook['Test Records']['E3'].value, 'SUBMITTED_CONFIRMED')
        self.assertTrue(workbook['Test Records']['G3'].hyperlink.target.startswith('http://127.0.0.1:'))
        self.assertGreater(workbook['Field Results'].max_row, 1)
        self.assertGreater(workbook['Detected Fields'].max_row, 1)
        self.assertGreater(workbook['Errors'].max_row, 1)
        test_class = type(self)
        test_class.server.shutdown()
        test_class.server_thread.join(timeout=5)
        test_class.server.server_close()
        test_class.server = make_server('127.0.0.1', test_class.port, application.app, threaded=True)
        test_class.server_thread = threading.Thread(target=test_class.server.serve_forever, daemon=True)
        test_class.server_thread.start()
        restarted_client = application.app.test_client()
        login_page = restarted_client.get('/login')
        login_token = self._csrf(login_page.data)
        restarted_client.post(
            '/login',
            data={'csrf_token': login_token, 'password': 'browser-local-password'},
        )
        restarted_history = restarted_client.get(
            '/history?test_name=Valid+Browser+Tester&status=SUBMITTED_CONFIRMED'
        )
        self.assertIn(b'Valid Browser Tester', restarted_history.data)

    def test_daily_limit_cannot_exceed_six_and_blocks_seventh_run(self):
        response = self.client.post(
            '/settings/daily-limit',
            data={'csrf_token': self.token, 'daily_limit': '7'},
            follow_redirects=True,
        )
        self.assertIn(b'Daily limit must be between 1 and 6.', response.data)
        with application.db() as con:
            con.execute(
                "UPDATE app_settings SET setting_value='6' WHERE setting_key='daily_limit'"
            )
        fields = self._open_and_detect()
        rules = self._rules(fields)
        for index in range(1, 7):
            response = self._post(
                '/browser/preview',
                {'name': f'Daily Test {index}', 'rules': rules},
            )
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        seventh = self._post('/browser/preview', {'name': 'Daily Test 7', 'rules': rules})
        self.assertEqual(seventh.status_code, 429)

    def test_unrecognized_success_page_is_not_confirmation(self):
        fields = self._open_and_detect()
        application.browser_worker.call(
            'set_page_content',
            '<form><button type="submit">Send</button></form><h1>Success</h1>',
        )
        result = application.browser_worker.call('submit_and_confirm')
        self.assertEqual(result['status'], 'SUBMISSION_UNVERIFIED')

    def test_dynamic_change_after_validation_blocks_submission(self):
        fields = self._open_and_detect()
        application.browser_worker.call('check_selector', '[name="initial_consent"]')
        application.browser_worker.call('check_selector', '[name="final_consent"]')
        preview = self._post('/browser/preview', {
            'name': 'Changed Page Run',
            'seed': 'stable',
            'rules': self._rules(fields),
        })
        self.assertEqual(preview.status_code, 200, preview.get_data(as_text=True))
        fill = self._post('/browser/fill', {})
        self.assertEqual(fill.get_json()['status'], 'VALIDATED')
        application.browser_worker.call('evaluate_page', """() => {
          const note = document.createElement('p');
          note.textContent = 'Dynamic page changed';
          document.body.append(note);
        }""")
        blocked = self._post('/browser/submit', {})
        self.assertEqual(blocked.status_code, 409)
        self.assertIn('Page navigation or dynamic changes', blocked.get_json()['error'])
        with application.db() as con:
            row = con.execute(
                'SELECT status FROM browser_runs WHERE id=?',
                (preview.get_json()['run_id'],),
            ).fetchone()
            local_count = con.execute('SELECT COUNT(*) FROM submissions').fetchone()[0]
        self.assertEqual(row['status'], 'VALIDATION_FAILED')
        self.assertEqual(local_count, 0)

    def test_dynamic_fields_iframe_and_custom_widget_notice(self):
        self._open_and_detect()
        application.browser_worker.call(
            'set_page_content',
            '<iframe title="nested form" srcdoc="<input name=&quot;hidden_from_main&quot;>"></iframe>'
            '<div role="combobox" aria-label="Custom choice">Open</div>'
            '<input type="password" name="login_value" value="DoNotExposeSecret">'
            '<input type="hidden" name="csrf_token" value="never-display">',
        )
        response = self._post('/browser/detect', {})
        self.assertEqual(response.status_code, 200)
        detected = response.get_json()
        self.assertEqual(detected['fields'], [])
        self.assertEqual(len(detected['frames']), 1)
        self.assertTrue(any(item['role'] == 'combobox' for item in detected['unsupported']))
        self.assertTrue(any(item['role'] == 'protected-field' for item in detected['unsupported']))
        self.assertNotIn('DoNotExposeSecret', response.get_data(as_text=True))
        self.assertNotIn('never-display', response.get_data(as_text=True))
        status = self.client.get('/browser/status').get_json()
        self.assertFalse(status['fields_stale'])

        application.browser_worker.call('evaluate_page', """() => {
          const input = document.createElement('input');
          input.name = 'dynamic_field';
          document.body.append(input);
        }""")
        status = self.client.get('/browser/status').get_json()
        self.assertTrue(status['fields_stale'])
        refreshed = self._post('/browser/detect', {}).get_json()
        self.assertEqual(len(refreshed['fields']), 1)
        self.assertEqual(refreshed['fields'][0]['name'], 'dynamic_field')

    def test_automation_dashboard_controls_render_in_browser(self):
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            page.goto(f'http://127.0.0.1:{self.port}/login')
            page.get_by_label('Dashboard password').fill('browser-local-password')
            page.get_by_role('button', name='Sign in').click()
            page.goto(f'http://127.0.0.1:{self.port}/automation')
            self.assertTrue(page.get_by_label('Target Website URL').is_visible())
            self.assertTrue(page.get_by_role('button', name='Open Website in Visible Browser').is_visible())
            self.assertTrue(page.get_by_role('button', name='Stop Browser').is_visible())
            self.assertIn('Daily run limit', page.locator('body').inner_text())
            queue_text = page.get_by_label('Enter names, one per line')
            queue_text.fill('Queue Tester A\nQueue Tester B')
            page.get_by_role('button', name='Save / replace queue').click()
            self.assertIn('Pending: 2', page.locator('#queue-counts').inner_text())
            page.get_by_role('button', name='Use name').first.click()
            self.assertEqual(page.locator('#test-name').input_value(), 'Queue Tester A')
            page.get_by_role('button', name='Mark skipped').first.click()
            self.assertIn('Skipped: 1', page.locator('#queue-counts').inner_text())
            page.reload()
            self.assertEqual(page.locator('#queue-list input').nth(0).input_value(), 'Queue Tester A')
            page.once('dialog', lambda dialog: dialog.accept())
            page.get_by_role('button', name='Delete local queue').click()
            self.assertIn('Pending: 0', page.locator('#queue-counts').inner_text())
            browser.close()
            playwright.stop()


if __name__ == '__main__':
    unittest.main()
